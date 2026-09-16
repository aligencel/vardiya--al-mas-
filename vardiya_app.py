import datetime
import io
import json
import os
import random
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Hastane Akıllı Vardiya Yönetim Sistemi",
    page_icon="🏥",
    layout="wide",
)

DATA_FILE = "hospital_shift_data.json"


# Veri Yükleme ve Kaydetme (Kalıcı Hafıza)
def load_data():
  if os.path.exists(DATA_FILE):
    try:
      with open(DATA_FILE, "r", encoding="utf-8") as f:
        return json.load(f)
    except:
      pass
  return {
      "Acil": {
          "personnel": [
              "MUSTAFA DİK",
              "TAYYİP DOĞAN",
              "HALİL VAR",
              "ADEM KÖMÜR",
              "SEDA YILMAZ",
              "FUNDA POYRAZ",
              "ENES ALTUN",
          ],
          "schedules": {},
      },
      "Yataklı Servis": {"personnel": ["Personel 1", "Personel 2"], "schedules": {}},
      "Palyatif": {"personnel": ["Personel A", "Personel B"], "schedules": {}},
  }


def save_data(data):
  with open(DATA_FILE, "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=4)


db = load_data()

tr_days_map = {
    "Monday": "PAZARTESİ",
    "Tuesday": "SALI",
    "Wednesday": "ÇARŞAMBA",
    "Thursday": "PERŞEMBE",
    "Friday": "CUMA",
    "Saturday": "CUMARTESİ",
    "Sunday": "PAZAR",
}

st.title("🏥 Hastane Birimleri Akıllı Vardiya Yönetim Sistemi")

# --- YAN MENÜ ---
with st.sidebar:
  st.header("📂 Birim Klasörleri")
  units = list(db.keys())
  selected_unit = st.selectbox("Çalışılacak Birimi Seçin", units)

  new_unit_name = st.text_input("➕ Yeni Birim Ekle (Örn: Yoğun Bakım)")
  if st.button("Birimi Oluştur"):
    if (
        new_unit_name
        and new_unit_name.strip()
        and new_unit_name not in db
    ):
      db[new_unit_name.strip()] = {"personnel": [], "schedules": {}}
      save_data(db)
      st.success(f"'{new_unit_name}' birimi eklendi!")
      st.rerun()

  st.markdown("---")
  st.header(f"👥 {selected_unit} - Personel Listesi")

  current_personnel = db[selected_unit]["personnel"]
  new_p_name = st.text_input("Yeni Personel Adı Soyadı")
  if st.button("Personeli Kaydet"):
    if new_p_name and new_p_name.strip().upper() not in current_personnel:
      db[selected_unit]["personnel"].append(new_p_name.strip().upper())
      save_data(db)
      st.success("Personel eklendi!")
      st.rerun()

  st.markdown("<b>Mevcut Personeller:</b>", unsafe_allow_html=True)
  to_remove = []
  # Hata çözümü için index (idx) eklendi
  for idx, p in enumerate(current_personnel):
    col_p1, col_p2 = st.columns([4, 1])
    col_p1.text(p)
    if col_p2.button("Sil", key=f"del_{selected_unit}_{idx}_{p}"):
      to_remove.append(p)

  if to_remove:
    db[selected_unit]["personnel"] = [
        p for p in current_personnel if p not in to_remove
    ]
    save_data(db)
    st.rerun()

  st.markdown("---")
  st.header("📅 Tarih ve Süre Ayarları")
  start_date = st.date_input(
      "Başlangıç Tarihi", value=datetime.date(2026, 9, 21)
  )

  duration_option = st.selectbox(
      "Döküm Süresi",
      [
          ("1 Hafta (7 Gün)", 7),
          ("2 Hafta (14 Gün)", 14),
          ("3 Hafta (21 Gün)", 21),
          ("4 Hafta (28 Gün)", 28),
      ],
      format_func=lambda x: x[0],
      index=2,
  )
  num_days = duration_option[1]

  st.markdown("---")
  generate_btn = st.button(
      "🎲 Bu Birim İçin Plan Oluştur", type="primary"
  )

# Tarih aralığını oluştur
days_info = []
current_d = start_date
for i in range(num_days):
  date_str = current_d.strftime("%d.%m.%Y")
  eng_day_name = current_d.strftime("%A")
  tr_day_name = tr_days_map.get(eng_day_name, eng_day_name)
  days_info.append(
      {"date_obj": current_d, "date_str": date_str, "day_name": tr_day_name}
  )
  current_d += datetime.timedelta(days=1)

personnel_list = db[selected_unit]["personnel"]
schedule_key = f"{start_date}_{num_days}"


# Tek bir haftayı kesin kurallarla çözen fonksiyon
def solve_single_week(personnel_list, prev_sunday_night=None):
  if prev_sunday_night is None:
    prev_sunday_night = []

  week_days = [
      "Pazartesi",
      "Salı",
      "Çarşamba",
      "Perşembe",
      "Cuma",
      "Cumartesi",
      "Pazar",
  ]

  for attempt in range(40000):
    week_data = {}
    gunduz_left = {p: 2 for p in personnel_list}
    gece_left = {p: 2 for p in personnel_list}

    consec_work = {p: 0 for p in personnel_list}
    consec_off = {p: 0 for p in personnel_list}
    consec_day = {p: 0 for p in personnel_list}
    consec_night = {p: 0 for p in personnel_list}
    last_shift = {p: None for p in personnel_list}

    success = True
    for d_idx, day in enumerate(week_days):
      forced = [
          p
          for p in personnel_list
          if (gunduz_left[p] + gece_left[p]) > 0 and consec_off[p] >= 2
      ]
      if len(forced) > 4:
        success = False
        break

      # Gece adayları (2 kişi)
      gece_cands = []
      for p in personnel_list:
        if gece_left[p] <= 0:
          continue
        if consec_work[p] >= 2:
          continue
        if consec_night[p] >= 2:
          continue
        gece_cands.append(p)

      if len(gece_cands) < 2:
        success = False
        break


      def score_g(p):
        s = 0
        if p in forced:
          s -= 100
        if last_shift[p] == "Gece":
          s += 10
        return s + random.uniform(0, 1)


      gece_cands.sort(key=score_g)
      gece_workers = gece_cands[:2]

      # Gündüz adayları (2 kişi)
      gunduz_cands = []
      for p in personnel_list:
        if gunduz_left[p] <= 0:
          continue
        if p in gece_workers:
          continue
        if consec_work[p] >= 2:
          continue
        if consec_day[p] >= 2:
          continue
        if last_shift[p] == "Gece":
          continue  # Gece çıkışı gündüze yasak
        if d_idx == 0 and p in prev_sunday_night:
          continue

        gunduz_cands.append(p)

      if len(gunduz_cands) < 2:
        success = False
        break


      def score_d(p):
        s = 0
        if p in forced:
          s -= 100
        return s + random.uniform(0, 1)


      gunduz_cands.sort(key=score_d)
      gunduz_workers = gunduz_cands[:2]

      today_workers = set(gece_workers + gunduz_workers)
      if not all(fw in today_workers for fw in forced):
        success = False
        break

      week_data[day] = {"Gündüz": gunduz_workers, "Gece": gece_workers}

      for p in personnel_list:
        if p in gunduz_workers:
          gunduz_left[p] -= 1
        if p in gece_workers:
          gece_left[p] -= 1

        if p in today_workers:
          consec_work[p] += 1
          consec_off[p] = 0
          if p in gece_workers:
            consec_day[p] = 0
            if last_shift[p] == "Gece":
              consec_night[p] += 1
            else:
              consec_night[p] = 1
            last_shift[p] = "Gece"
          else:
            consec_night[p] = 0
            if last_shift[p] == "Gündüz":
              consec_day[p] += 1
            else:
              consec_day[p] = 1
            last_shift[p] = "Gündüz"
        else:
          consec_work[p] = 0
          consec_off[p] += 1
          consec_day[p] = 0
          consec_night[p] = 0
          last_shift[p] = None

    if (
        success
        and all(gunduz_left[p] == 0 for p in personnel_list)
        and all(gece_left[p] == 0 for p in personnel_list)
    ):
      return week_data

  return None


# Plan üretme algoritması
if generate_btn:
  if len(personnel_list) < 4:
    st.error(
        "Vardiya planı oluşturabilmek için bu birimde en az 4 personel"
        " olmalıdır!"
    )
  else:
    num_weeks = num_days // 7
    full_schedule = {}
    prev_sunday_night = []
    generation_success = True

    for w in range(num_weeks):
      week_result = solve_single_week(personnel_list, prev_sunday_night)
      if week_result is None:
        generation_success = False
        break

      week_days_slice = days_info[w * 7 : (w + 1) * 7]
      week_day_names = [
          "Pazartesi",
          "Salı",
          "Çarşamba",
          "Perşembe",
          "Cuma",
          "Cumartesi",
          "Pazar",
      ]

      for idx, d_info in enumerate(week_days_slice):
        d_name = week_day_names[idx]
        d_str = d_info["date_str"]
        full_schedule[d_str] = week_result[d_name]

      prev_sunday_night = week_result["Pazar"]["Gece"]

    if not generation_success:
      st.error(
          "Seçilen tarih aralığı için tüm haftalık dengeler kurallara uygun"
          " şekilde kurulamadı. Lütfen tekrar deneyin."
      )
    else:
      matrix_rows = []
      for p in personnel_list:
        row = {"Personel": p}
        total_days = 0
        g_count = 0
        n_count = 0
        for day_info in days_info:
          d_str = day_info["date_str"]
          if p in full_schedule[d_str]["Gündüz"]:
            row[d_str] = "08:00-20:00"
            total_days += 1
            g_count += 1
          elif p in full_schedule[d_str]["Gece"]:
            row[d_str] = "20:00-08:00"
            total_days += 1
            n_count += 1
          else:
            row[d_str] = "HT"
        row["Gündüz"] = g_count
        row["Gece"] = n_count
        row["Toplam Gün"] = total_days
        matrix_rows.append(row)

      db[selected_unit]["schedules"][schedule_key] = matrix_rows
      save_data(db)
      st.success(
          f"'{selected_unit}' birimi için her haftası kesin 2 Gündüz / 2 Gece"
          " dengeli plan başarıyla oluşturuldu!"
      )

# --- ANA EKRANDA TABLOYU GÖSTERME ---
st.subheader(
    f"📅 {selected_unit} - Vardiya Matris Tablosu ({num_days} Günlük)"
)

saved_schedules = db[selected_unit]["schedules"]
if schedule_key in saved_schedules:
  df = pd.DataFrame(saved_schedules[schedule_key])

  html_code = f"""
    <style>
    .shift-container {{ width: 100%; overflow-x: auto; }}
    .shift-table {{ width: 100%; border-collapse: collapse; font-family: Arial, sans-serif; font-size: 12px; box-shadow: 0 2px 5px rgba(0,0,0,0.1); }}
    .shift-table th, .shift-table td {{ border: 1px solid #b0b0b0; padding: 8px 4px; text-align: center; white-space: nowrap; }}
    .shift-table th {{ background-color: #388E3C; color: white; font-weight: bold; font-size: 11px; }}
    .shift-table td.name-col {{ background-color: #f5f5f5; font-weight: bold; text-align: left; color: #000; padding-left: 10px; }}
    .day-shift {{ background-color: #FFC107 !important; color: #000000 !important; font-weight: bold; }}
    .night-shift {{ background-color: #42A5F5 !important; color: #000000 !important; font-weight: bold; }}
    .off-shift {{ background-color: #E0E0E0 !important; color: #555555 !important; font-weight: bold; }}
    </style>
    <div class="shift-container">
    <table class="shift-table">
        <thead>
            <tr>
                <th style="background-color: #2E7D32; text-align: left; padding-left: 10px; vertical-align: middle; font-size: 14px;">{selected_unit.upper()}</th>
    """
  for day_info in days_info:
    html_code += f"<th><span style='font-size: 10px; opacity: 0.9;'>{day_info['date_str']}</span><br><span style='font-size: 12px;'>{day_info['day_name']}</span></th>"
  html_code += (
      "<th>GÜNDÜZ</th><th>GECE</th><th>TOPLAM</th></tr></thead><tbody>"
  )

  for idx, row in df.iterrows():
    html_code += f"<tr><td class='name-col'>{row['Personel']}</td>"
    for day_info in days_info:
      d_str = day_info["date_str"]
      val = row.get(d_str, "HT")
      if val == "08:00-20:00":
        html_code += f"<td class='day-shift'>{val}</td>"
      elif val == "20:00-08:00":
        html_code += f"<td class='night-shift'>{val}</td>"
      else:
        html_code += f"<td class='off-shift'>{val}</td>"
    html_code += f"<td style='background-color: #f9f9f9;'><b>{row['Gündüz']}</b></td><td style='background-color: #f9f9f9;'><b>{row['Gece']}</b></td><td style='background-color: #f9f9f9;'><b>{row['Toplam Gün']}</b></td></tr>"

  html_code += "</tbody></table></div>"
  st.markdown(html_code, unsafe_allow_html=True)

  st.markdown("---")

  output = io.BytesIO()
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = selected_unit[:31]

  green_header_fill = PatternFill(
      start_color="388E3C", end_color="388E3C", fill_type="solid"
  )
  yellow_fill = PatternFill(
      start_color="FFC107", end_color="FFC107", fill_type="solid"
  )
  blue_fill = PatternFill(
      start_color="42A5F5", end_color="42A5F5", fill_type="solid"
  )
  gray_fill = PatternFill(
      start_color="E0E0E0", end_color="E0E0E0", fill_type="solid"
  )
  light_gray_fill = PatternFill(
      start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"
  )
  summary_fill = PatternFill(
      start_color="F9F9F9", end_color="F9F9F9", fill_type="solid"
  )

  white_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
  bold_font = Font(name="Arial", size=10, bold=True, color="000000")
  regular_font = Font(name="Arial", size=10, color="555555")

  thin_border = Border(
      left=Side(style="thin", color="B0B0B0"),
      right=Side(style="thin", color="B0B0B0"),
      top=Side(style="thin", color="B0B0B0"),
      bottom=Side(style="thin", color="B0B0B0"),
  )

  headers = list(df.columns)
  ws.append(headers)
  ws.row_dimensions[1].height = 25

  for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = green_header_fill
    cell.font = white_font
    cell.alignment = Alignment(
        horizontal="left" if col_num == 1 else "center",
        vertical="center",
        wrap_text=True,
    )
    cell.border = thin_border

  for row_idx, row_data in enumerate(df.values, 2):
    ws.row_dimensions[row_idx].height = 22
    for col_num, val in enumerate(row_data, 1):
      cell = ws.cell(row=row_idx, column=col_num, value=val)
      cell.border = thin_border
      cell.alignment = Alignment(horizontal="center", vertical="center")

      if col_num == 1:
        cell.fill = light_gray_fill
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
      elif val == "08:00-20:00":
        cell.fill = yellow_fill
        cell.font = bold_font
      elif val == "20:00-08:00":
        cell.fill = blue_fill
        cell.font = bold_font
      elif val == "HT":
        cell.fill, cell.font = gray_fill, regular_font
      else:
        if col_num > len(headers) - 3:
          cell.fill = summary_fill
          cell.font = bold_font

  wb.save(output)
  excel_data = output.getvalue()

  col_dl1, col_dl2 = st.columns([1, 3])
  with col_dl1:
    st.download_button(
        label="📥 Renkli Tabloyu Excel Olarak İndir (.xlsx)",
        data=excel_data,
        file_name=f"{selected_unit}_Vardiya_Plani.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
  with col_dl2:
    st.success("Plan başarıyla oluşturuldu ve Excel için hazırlandı.")

else:
  st.warning(
      "Bu birim ve tarih aralığı için henüz kaydedilmiş bir vardiya planı"
      " bulunmuyor. Sol menüden **'Bu Birim İçin Plan Oluştur'** butonuna"
      " basarak planı üretebilirsiniz."
  )